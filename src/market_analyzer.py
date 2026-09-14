"""Market analysis, zip code scoring, and Market Finder reports.

Analyzes county-level data to identify target zip codes for investment.
Scores each zip by distress density, property values, equity, and competition.

Data source: the CSVs already sitting in OUTPUT_DIR, and nothing else. This
module makes no network calls at analysis time. Property values, equity and
tax-delinquency dollars are read from whatever enrichment a previous run wrote
into those CSVs — the Zillow and county-tax APIs named in earlier versions of
this docstring are upstream producers, not runtime dependencies, and
`zillow_market_api` is never imported here.

Counties must be passed explicitly for any market outside Knox/Blount; the
known-zip backfill below only covers those two.

Usage:
  python src/main.py market-analysis --counties Knox,Blount
  python src/main.py market-analysis --counties Essex,Middlesex,Somerset,Union
  python src/main.py market-analysis --counties Knox --zip-codes 37918,37919,37920
"""

import csv
import glob
import logging
import os
import random
import re
import statistics
import time
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import config

logger = logging.getLogger(__name__)

# ── Knox/Blount county zip codes ──────────────────────────────────────
KNOX_ZIPS = [
    "37901", "37902", "37909", "37912", "37914", "37915", "37916", "37917",
    "37918", "37919", "37920", "37921", "37922", "37923", "37924", "37931",
    "37932", "37934", "37938",
]
BLOUNT_ZIPS = [
    "37801", "37803", "37804", "37853", "37882", "37886",
]
COUNTY_ZIPS = {
    "knox": KNOX_ZIPS,
    "blount": BLOUNT_ZIPS,
}

# ── ZIP validity spans ────────────────────────────────────────────────
# A ZIP that cannot belong to the state its county sits in is address-parse
# junk, not a market signal: the corpus carries rows filed under Middlesex
# County NJ with ZIPs in New York (12590) and North Carolina (28879). The
# `state` column does not catch these — it is stamped "NJ" on every row rather
# than parsed — so the ZIP itself is the only available check.
#
# These spans are deliberately state-wide, not county-tight. They reject only
# provably-impossible ZIPs and assert nothing about which ZIPs a county
# actually contains, so an incomplete list here can never discard real data.
# Residual case they do not catch: a valid in-state ZIP filed under the wrong
# county (08210, Cape May Court House, filed under Middlesex). Catching that
# needs per-county ZIP sets, which are not asserted here.
COUNTY_STATE = {
    "knox": "TN", "blount": "TN",
    "essex": "NJ", "middlesex": "NJ", "somerset": "NJ", "union": "NJ",
}
STATE_ZIP_SPANS = {
    "NJ": (7001, 8989),
    "TN": (37010, 38589),
}

_ADDR_NOISE = re.compile(r"[^A-Z0-9 ]")
_ADDR_SPACE = re.compile(r"\s+")


def _norm_address(value: str) -> str:
    """Collapse an address into a comparison key: upper, alphanumeric, single-spaced."""
    return _ADDR_SPACE.sub(" ", _ADDR_NOISE.sub(" ", (value or "").upper())).strip()

# ── Scoring weights ───────────────────────────────────────────────────
WEIGHT_DISTRESS_DENSITY = 0.30   # Foreclosure/tax sale/probate count per zip
WEIGHT_MEDIAN_VALUE = 0.20       # Lower median = better for investors
WEIGHT_EQUITY_AVG = 0.15         # Higher equity = more room for deals
WEIGHT_TAX_DELINQUENCY = 0.15   # More delinquency = more distress
WEIGHT_COMPETITION = 0.10        # Lower investor activity = less competition
WEIGHT_DOM_AVG = 0.10            # Higher DOM = more negotiating leverage

# ── Data structures ───────────────────────────────────────────────────


@dataclass
class ZipProfile:
    """Profile for a single zip code."""
    zip_code: str = ""
    county: str = ""
    # Notice/distress data from our scraped records
    total_notices: int = 0
    foreclosure_count: int = 0
    tax_sale_count: int = 0
    tax_delinquent_count: int = 0
    probate_count: int = 0
    eviction_count: int = 0
    code_violation_count: int = 0
    sheriff_sale_count: int = 0
    # Property data (aggregated from Zillow)
    median_value: float = 0.0
    avg_equity_pct: float = 0.0
    avg_dom: float = 0.0
    property_count: int = 0
    # Tax delinquency data
    avg_tax_delinquent_amount: float = 0.0
    tax_delinquent_property_count: int = 0
    # Competition (investor activity)
    investor_purchase_count: int = 0
    competition_ratio: float = 0.0  # investor purchases / total sales
    # Raw samples, retained so the aggregates above are true statistics rather
    # than order-dependent running averages. Not written to the workbook.
    value_samples: list = field(default_factory=list)
    equity_samples: list = field(default_factory=list)
    tax_samples: list = field(default_factory=list)
    # Calculated score
    score: float = 0.0
    rank: int = 0
    grade: str = ""  # A, B, C, D


@dataclass
class CorpusStats:
    """What one `_load_notice_data` pass read, and what it discarded.

    Carried into the workbook so a grade can be traced back to the corpus that
    produced it — the analyzer reads a scratch directory nobody curates, and
    the deductions below are the difference between a defensible score and a
    duplicate-file artifact.
    """
    files_read: list = field(default_factory=list)
    rows_scanned: int = 0
    rows_kept: int = 0
    duplicates_dropped: int = 0
    malformed_zips_dropped: int = 0
    out_of_state_zips_dropped: int = 0
    # Deduplicated reasons, not one per row — a junk row in a triplicated file
    # would otherwise be reported three times.
    rejected_zips: dict = field(default_factory=dict)
    rows_without_address: int = 0


@dataclass
class MarketReport:
    """Complete market analysis report."""
    county: str = ""
    analysis_date: str = ""
    total_zips: int = 0
    zip_profiles: list = field(default_factory=list)
    top_zips: list = field(default_factory=list)
    total_notices: int = 0
    avg_median_value: float = 0.0
    avg_equity: float = 0.0
    total_distress: int = 0
    corpus: CorpusStats = field(default_factory=CorpusStats)
    factors: list = field(default_factory=list)


# ── Data aggregation from our CSVs ───────────────────────────────────

def _load_notice_data(counties: list[str] | None = None,
                      stats: "CorpusStats | None" = None) -> dict[str, ZipProfile]:
    """Aggregate notice data from our output CSVs by zip code.

    Reads every CSV in OUTPUT_DIR in sorted order and counts each distinct
    (address, zip, notice_type) once. Both halves of that matter:

    **Sorted, not glob order.** `Path.glob` returns filesystem order, which is
    arbitrary and OS-dependent. The per-zip aggregates below are computed from
    the rows in the order they arrive, so glob order made the output
    irreproducible: the same corpus on the same machine scored a top zip's
    equity anywhere from 52% to 92% depending only on which file the OS handed
    back first, and adding an unrelated file to OUTPUT_DIR silently reordered
    the rankings.

    **Deduplicated.** OUTPUT_DIR is an append-only scratch area, not a curated
    corpus: it holds reimport chains and held-for-cleaning copies of the same
    run. Counting rows counted one 1,187-record probate backfill three times
    and made the leading zip in the corpus substantially a duplicate-file
    artifact.

    A property contributes at most one notice per type. `Date Added` is empty
    on every row of the present corpus, so it cannot separate a genuine
    re-filing from a reimport and is deliberately not part of the key; a row
    with no address at all cannot be keyed and is kept but counted, so it
    shows up in the provenance sheet rather than vanishing.
    """
    profiles = defaultdict(lambda: ZipProfile())
    stats = stats if stats is not None else CorpusStats()

    csv_files = sorted(config.OUTPUT_DIR.glob("*.csv"))
    if not csv_files:
        logger.warning("No CSV files found in %s", config.OUTPUT_DIR)
        return dict(profiles)

    county_filter = {c.lower() for c in counties} if counties else None
    seen: set[tuple[str, str, str]] = set()

    type_map = {
        "foreclosure": "foreclosure_count",
        "tax_sale": "tax_sale_count",
        "tax_delinquent": "tax_delinquent_count",
        "probate": "probate_count",
        "eviction": "eviction_count",
        "code_violation": "code_violation_count",
        "sheriff_sale": "sheriff_sale_count",
    }

    for csv_path in csv_files:
        stats.files_read.append(csv_path.name)
        try:
            with open(csv_path, "r", encoding="utf-8-sig") as f:
                reader = csv.DictReader(f)
                for row in reader:
                    stats.rows_scanned += 1

                    zip_code = (row.get("zip") or row.get("ZIP") or "").strip()[:5]
                    if not zip_code or not zip_code.isdigit():
                        continue

                    county = (row.get("county") or row.get("County") or "").strip().lower()
                    if county_filter and county not in county_filter:
                        continue

                    city = (row.get("city") or row.get("City") or "").strip()
                    if len(zip_code) != 5:
                        stats.malformed_zips_dropped += 1
                        stats.rejected_zips.setdefault(
                            zip_code,
                            f"{zip_code} ({city}) under {county.title()} County — "
                            f"not a 5-digit ZIP",
                        )
                        continue

                    span = STATE_ZIP_SPANS.get(COUNTY_STATE.get(county, ""))
                    if span and not span[0] <= int(zip_code) <= span[1]:
                        stats.out_of_state_zips_dropped += 1
                        stats.rejected_zips.setdefault(
                            zip_code,
                            f"{zip_code} ({city}) under {county.title()} County — "
                            f"outside the state ZIP span",
                        )
                        continue

                    notice_type = (row.get("notice_type") or row.get("Notice Type") or "").strip().lower()

                    address = _norm_address(row.get("address") or row.get("property_address"))
                    if address:
                        key = (address, zip_code, notice_type)
                        if key in seen:
                            stats.duplicates_dropped += 1
                            continue
                        seen.add(key)
                    else:
                        stats.rows_without_address += 1

                    stats.rows_kept += 1
                    profile = profiles[zip_code]
                    profile.zip_code = zip_code
                    profile.county = county.title()
                    profile.total_notices += 1

                    attr = type_map.get(notice_type)
                    if attr:
                        setattr(profile, attr, getattr(profile, attr) + 1)

                    # Property data is collected as raw samples and reduced once,
                    # after the whole corpus is in — see _finalize_profiles.
                    est_val = row.get("estimated_value") or row.get("Estimated Value") or ""
                    if est_val:
                        try:
                            val = float(est_val.replace(",", "").replace("$", ""))
                            if val > 0:
                                profile.value_samples.append(val)
                        except ValueError:
                            pass

                    equity_pct = row.get("equity_percent") or row.get("Equity Percentage") or ""
                    if equity_pct:
                        try:
                            profile.equity_samples.append(float(equity_pct.replace("%", "")))
                        except ValueError:
                            pass

                    tax_amt = row.get("tax_delinquent_amount") or ""
                    if tax_amt:
                        try:
                            amt = float(tax_amt.replace(",", "").replace("$", ""))
                            if amt > 0:
                                profile.tax_samples.append(amt)
                        except ValueError:
                            pass

        except Exception as e:
            logger.debug("Error reading %s: %s", csv_path, e)

    _finalize_profiles(profiles)

    logger.info(
        "Corpus: %d files, %d rows scanned, %d kept (%d duplicate property-notices, "
        "%d malformed ZIPs, %d out-of-state ZIPs dropped; %d rows had no address to key on)",
        len(stats.files_read), stats.rows_scanned, stats.rows_kept,
        stats.duplicates_dropped, stats.malformed_zips_dropped,
        stats.out_of_state_zips_dropped, stats.rows_without_address,
    )
    for reason in stats.rejected_zips.values():
        logger.warning("Rejected ZIP: %s", reason)

    return dict(profiles)


def _finalize_profiles(profiles: dict) -> None:
    """Reduce each profile's raw samples to true statistics, in place.

    Done once at the end rather than incrementally because the incremental
    forms were both wrong and order-dependent: `median_value` was a running
    mean carrying a "median" label, and equity used `(avg + pct) / 2`, which
    weights the most recent row 50%, the one before it 25%, and so on. Equity
    is 15% of the composite, and that recency weighting moved the top zip by
    ~6 points depending on file order alone.
    """
    for profile in profiles.values():
        profile.property_count = len(profile.value_samples)
        profile.median_value = (
            statistics.median(profile.value_samples) if profile.value_samples else 0.0
        )
        profile.avg_equity_pct = (
            statistics.fmean(profile.equity_samples) if profile.equity_samples else 0.0
        )
        profile.tax_delinquent_property_count = len(profile.tax_samples)
        profile.avg_tax_delinquent_amount = (
            statistics.fmean(profile.tax_samples) if profile.tax_samples else 0.0
        )


# ── Scoring engine ────────────────────────────────────────────────────

def _normalize(values: list[float], higher_is_better: bool = True) -> list[float]:
    """Normalize values to 0-100 scale."""
    if not values:
        return []
    mn, mx = min(values), max(values)
    if mn == mx:
        return [50.0] * len(values)
    normed = [(v - mn) / (mx - mn) * 100 for v in values]
    if not higher_is_better:
        normed = [100 - n for n in normed]
    return normed


# (label, weight, ZipProfile attribute, higher_is_better)
_SCORING_FACTORS = (
    ("Distress density", WEIGHT_DISTRESS_DENSITY, "total_notices", True),
    ("Median value", WEIGHT_MEDIAN_VALUE, "median_value", False),
    ("Equity", WEIGHT_EQUITY_AVG, "avg_equity_pct", True),
    ("Tax delinquency", WEIGHT_TAX_DELINQUENCY, "avg_tax_delinquent_amount", True),
    ("Competition", WEIGHT_COMPETITION, "competition_ratio", False),
    ("Days on market", WEIGHT_DOM_AVG, "avg_dom", True),
)


def score_zip_codes(profiles: dict[str, ZipProfile],
                    factor_log: list | None = None) -> list[ZipProfile]:
    """Score and rank zip codes for investment potential.

    A factor identical across every zip carries no information. `_normalize`
    returns a flat 50.0 for it, and a flat 50 is not neutral: it pulls every
    score toward the midpoint. Three of the six factors are unpopulated in
    practice — `competition_ratio` and `avg_dom` are declared and read but
    never assigned anywhere in the repo, and `avg_tax_delinquent_amount` has
    one population path that is gated to a single county — so 0.35 of the
    weight was returning a constant. The resulting transform,
    `score = 0.65 * live + 17.5`, squeezed every possible score into
    [17.5, 82.5]: an A (>= 75) required 88.5 on the live factors, so no zip in
    the corpus could earn one, and nothing could fall below a floor of 17.5.

    So a constant factor is dropped and the surviving weights renormalized to
    1.0. Scores then span the full 0-100 and a grade means what the grade
    table says it means. Dropping is computed, not hardcoded: if a dead factor
    is ever populated it re-enters scoring on its own, with no code change.
    Every drop is logged, and `factor_log`, when supplied, receives one record
    per factor for the workbook's provenance sheet.

    Ranking is unaffected by this change — a constant term is monotonic — but
    the grades are not, and neither is the defensibility of a grade.
    """
    if not profiles:
        return []

    zips = list(profiles.values())

    live: list[tuple[str, float, list[float]]] = []
    dropped: list[tuple[str, float, float]] = []
    for label, weight, attr, higher_is_better in _SCORING_FACTORS:
        values = [float(getattr(z, attr)) for z in zips]
        if min(values) == max(values):
            dropped.append((label, weight, values[0]))
        else:
            live.append((label, weight, _normalize(values, higher_is_better=higher_is_better)))

    if not live:
        raise ValueError(
            f"No scoring factor varies across the {len(zips)} zips supplied — every "
            f"factor is constant, so there is nothing to rank. Check that the corpus "
            f"actually loaded before trusting any output."
        )

    total_weight = sum(weight for _, weight, _ in live)
    for label, weight, value in dropped:
        logger.warning(
            "Scoring factor %r is constant at %g across all %d zips — dropped from "
            "scoring (weight %.2f). A flat factor compresses every score toward the "
            "midpoint rather than leaving them unchanged.",
            label, value, len(zips), weight,
        )
    logger.info(
        "Scoring on %d of %d factors; surviving weights renormalized from %.2f to 1.00",
        len(live), len(_SCORING_FACTORS), total_weight,
    )

    for i, z in enumerate(zips):
        z.score = sum(
            normed[i] * weight / total_weight for _, weight, normed in live
        )

    if factor_log is not None:
        factor_log.extend(
            {"factor": label, "status": "live", "weight_declared": weight,
             "weight_applied": weight / total_weight, "note": ""}
            for label, weight, _ in live
        )
        factor_log.extend(
            {"factor": label, "status": "DROPPED", "weight_declared": weight,
             "weight_applied": 0.0,
             "note": f"constant at {value:g} across all zips — no discriminating power"}
            for label, weight, value in dropped
        )

    # Sort by score descending
    zips.sort(key=lambda z: z.score, reverse=True)

    # Assign ranks and grades
    for i, z in enumerate(zips):
        z.rank = i + 1
        pct = z.score
        if pct >= 75:
            z.grade = "A"
        elif pct >= 55:
            z.grade = "B"
        elif pct >= 35:
            z.grade = "C"
        else:
            z.grade = "D"

    return zips


# ── Budget allocation ─────────────────────────────────────────────────

def _allocate_budget(zips: list[ZipProfile], monthly_budget: float = 5000.0,
                     max_zips: int = 5) -> list[tuple[str, float, str]]:
    """Allocate marketing budget across top zip codes by score weight."""
    top = zips[:max_zips]
    if not top:
        return []

    total_score = sum(z.score for z in top)
    if total_score == 0:
        equal = monthly_budget / len(top)
        return [(z.zip_code, round(equal), z.grade) for z in top]

    return [
        (z.zip_code, round(monthly_budget * z.score / total_score), z.grade)
        for z in top
    ]


# ── Excel report ──────────────────────────────────────────────────────

_HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
_HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
_TITLE_FONT = Font(name="Calibri", bold=True, size=16, color="2F5496")
_SUBTITLE_FONT = Font(name="Calibri", bold=True, size=12, color="333333")
_LABEL_FONT = Font(name="Calibri", size=11, color="555555")
_VALUE_FONT = Font(name="Calibri", bold=True, size=13, color="222222")
_GRADE_COLORS = {
    "A": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    "B": PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid"),
    "C": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
    "D": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
}
_THIN_BORDER = Border(bottom=Side(style="thin", color="D9D9D9"))
_MONEY_FMT = '#,##0'


def _write_headers(ws, row, headers):
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN


def _auto_widths(ws, min_w=12, max_w=30):
    for col in ws.columns:
        mx = max((len(str(c.value or "")) for c in col), default=0)
        ws.column_dimensions[col[0].column_letter].width = min(max(mx + 2, min_w), max_w)


def generate_market_report(report: MarketReport, budget: list[tuple],
                           output_path: str = "") -> str:
    """Generate a 6-tab Market Finder Excel workbook."""
    wb = Workbook()

    # ── Tab 1: Executive Summary ──────────────────────────────────
    ws = wb.active
    ws.title = "Executive Summary"
    ws.cell(row=1, column=1, value="Market Analysis Report").font = _TITLE_FONT
    ws.cell(row=2, column=1, value=f"County: {report.county}").font = _SUBTITLE_FONT
    ws.cell(row=3, column=1, value=f"Date: {report.analysis_date}").font = _LABEL_FONT

    ws.cell(row=5, column=1, value="Top 10 Target Zip Codes").font = _SUBTITLE_FONT
    _write_headers(ws, 6, ["Rank", "ZIP", "County", "Grade", "Score",
                            "Total Notices", "Median Value", "Avg Equity %"])
    for i, z in enumerate(report.top_zips[:10], 7):
        ws.cell(row=i, column=1, value=z.rank)
        ws.cell(row=i, column=2, value=z.zip_code)
        ws.cell(row=i, column=3, value=z.county)
        grade_cell = ws.cell(row=i, column=4, value=z.grade)
        grade_cell.fill = _GRADE_COLORS.get(z.grade, PatternFill())
        ws.cell(row=i, column=5, value=round(z.score, 1))
        ws.cell(row=i, column=6, value=z.total_notices)
        ws.cell(row=i, column=7, value=round(z.median_value)).number_format = _MONEY_FMT
        ws.cell(row=i, column=8, value=f"{z.avg_equity_pct:.0f}%" if z.avg_equity_pct else "")
        for c in range(1, 9):
            ws.cell(row=i, column=c).border = _THIN_BORDER

    summary_row = 18
    summary_data = [
        ("Total Zip Codes Analyzed", str(report.total_zips)),
        ("Total Distress Notices", str(report.total_notices)),
        ("Avg Median Home Value", f"${report.avg_median_value:,.0f}" if report.avg_median_value else "N/A"),
        ("Avg Equity", f"{report.avg_equity:.0f}%" if report.avg_equity else "N/A"),
    ]
    for label, value in summary_data:
        ws.cell(row=summary_row, column=1, value=label).font = _LABEL_FONT
        ws.cell(row=summary_row, column=2, value=value).font = _VALUE_FONT
        summary_row += 1

    _auto_widths(ws)

    # ── Tab 2: Zip Scorecard ──────────────────────────────────────
    ws2 = wb.create_sheet("Zip Scorecard")
    ws2.cell(row=1, column=1, value="All Zip Codes — Scored & Ranked").font = _TITLE_FONT
    headers = ["Rank", "ZIP", "County", "Grade", "Score", "Total Notices",
               "Foreclosures", "Sheriff Sales", "Tax Sales", "Tax Delinquent",
               "Probate", "Evictions", "Code Violations", "Median Value",
               "Avg Equity %", "Avg Tax Delinquent $", "Properties Analyzed"]
    _write_headers(ws2, 3, headers)
    for i, z in enumerate(report.zip_profiles, 4):
        vals = [z.rank, z.zip_code, z.county, z.grade, round(z.score, 1),
                z.total_notices, z.foreclosure_count, z.sheriff_sale_count,
                z.tax_sale_count, z.tax_delinquent_count, z.probate_count,
                z.eviction_count, z.code_violation_count, round(z.median_value),
                f"{z.avg_equity_pct:.0f}%" if z.avg_equity_pct else "",
                round(z.avg_tax_delinquent_amount), z.property_count]
        for col, val in enumerate(vals, 1):
            cell = ws2.cell(row=i, column=col, value=val)
            if col == 4:
                cell.fill = _GRADE_COLORS.get(val, PatternFill())
            if col == 14:
                cell.number_format = _MONEY_FMT
            cell.border = _THIN_BORDER
    _auto_widths(ws2)

    # ── Tab 3: Distress Density ───────────────────────────────────
    ws3 = wb.create_sheet("Distress Density")
    ws3.cell(row=1, column=1, value="Notice Distribution by Type & ZIP").font = _TITLE_FONT
    ws3.cell(row=2, column=1,
             value="Raw notice counts, deduplicated to one notice per property per type. "
                   "Not yet a density — see the Method & Provenance tab.").font = _LABEL_FONT
    _write_headers(ws3, 3, ["ZIP", "County", "Foreclosure", "Sheriff Sale", "Tax Sale",
                             "Tax Delinquent", "Probate", "Eviction", "Code Violation",
                             "TOTAL"])
    for i, z in enumerate(sorted(report.zip_profiles, key=lambda x: x.total_notices, reverse=True), 4):
        vals = [z.zip_code, z.county, z.foreclosure_count, z.sheriff_sale_count,
                z.tax_sale_count, z.tax_delinquent_count, z.probate_count,
                z.eviction_count, z.code_violation_count, z.total_notices]
        for col, val in enumerate(vals, 1):
            ws3.cell(row=i, column=col, value=val).border = _THIN_BORDER
    _auto_widths(ws3)

    # ── Tab 4: Competition Map ────────────────────────────────────
    ws4 = wb.create_sheet("Competition Map")
    ws4.cell(row=1, column=1, value="Investor Activity by ZIP").font = _TITLE_FONT
    ws4.cell(row=2, column=1,
             value="Lower competition ratio = less investor activity = better opportunity").font = _LABEL_FONT
    _write_headers(ws4, 4, ["ZIP", "County", "Investor Purchases", "Competition Ratio", "Grade"])
    for i, z in enumerate(sorted(report.zip_profiles, key=lambda x: x.competition_ratio), 5):
        ws4.cell(row=i, column=1, value=z.zip_code)
        ws4.cell(row=i, column=2, value=z.county)
        ws4.cell(row=i, column=3, value=z.investor_purchase_count)
        ws4.cell(row=i, column=4, value=f"{z.competition_ratio:.1%}" if z.competition_ratio else "N/A")
        ws4.cell(row=i, column=5, value=z.grade).fill = _GRADE_COLORS.get(z.grade, PatternFill())
        for c in range(1, 6):
            ws4.cell(row=i, column=c).border = _THIN_BORDER
    _auto_widths(ws4)

    # ── Tab 5: Budget Recommendations ─────────────────────────────
    ws5 = wb.create_sheet("Recommendations")
    ws5.cell(row=1, column=1, value="Marketing Budget Allocation").font = _TITLE_FONT
    ws5.cell(row=2, column=1, value="Budget weighted by zip score — higher-scoring zips get more spend").font = _LABEL_FONT

    _write_headers(ws5, 4, ["ZIP", "Monthly Budget", "Grade", "% of Total"])
    total_budget = sum(b[1] for b in budget) if budget else 1
    for i, (zip_code, amount, grade) in enumerate(budget, 5):
        ws5.cell(row=i, column=1, value=zip_code)
        ws5.cell(row=i, column=2, value=amount).number_format = _MONEY_FMT
        ws5.cell(row=i, column=3, value=grade).fill = _GRADE_COLORS.get(grade, PatternFill())
        ws5.cell(row=i, column=4, value=f"{amount / total_budget:.0%}")
        for c in range(1, 5):
            ws5.cell(row=i, column=c).border = _THIN_BORDER
    _auto_widths(ws5)

    # ── Tab 6: Method & Provenance ────────────────────────────────
    # The analyzer reads an uncurated scratch directory and scores on whichever
    # factors happen to be populated. Without this tab a reader cannot tell a
    # real signal from a duplicate-file artifact, and no grade here is
    # defensible. Everything the score depends on that is not a notice count
    # belongs on this sheet.
    ws6 = wb.create_sheet("Method & Provenance")
    ws6.cell(row=1, column=1, value="How These Scores Were Produced").font = _TITLE_FONT

    r = 3
    ws6.cell(row=r, column=1, value="Corpus").font = _SUBTITLE_FONT
    r += 1
    c = report.corpus
    for label, value in [
        ("CSV files read", str(len(c.files_read))),
        ("Rows scanned", f"{c.rows_scanned:,}"),
        ("Rows counted", f"{c.rows_kept:,}"),
        ("Duplicate property-notices dropped", f"{c.duplicates_dropped:,}"),
        ("Malformed ZIPs dropped", str(c.malformed_zips_dropped)),
        ("Out-of-state ZIPs dropped", str(c.out_of_state_zips_dropped)),
        ("Rows with no address to key on", str(c.rows_without_address)),
    ]:
        ws6.cell(row=r, column=1, value=label).font = _LABEL_FONT
        ws6.cell(row=r, column=2, value=value).font = _VALUE_FONT
        r += 1

    r += 1
    ws6.cell(row=r, column=1,
             value="A property is counted once per notice type. The output directory is an "
                   "append-only scratch area holding reimport chains and held-for-cleaning "
                   "copies of the same run, so counting rows counts some records several "
                   "times over.").font = _LABEL_FONT
    r += 2

    ws6.cell(row=r, column=1, value="Scoring factors").font = _SUBTITLE_FONT
    r += 1
    _write_headers(ws6, r, ["Factor", "Status", "Declared Weight", "Applied Weight", "Note"])
    r += 1
    for f in report.factors:
        ws6.cell(row=r, column=1, value=f["factor"])
        status = ws6.cell(row=r, column=2, value=f["status"])
        if f["status"] == "DROPPED":
            status.fill = _GRADE_COLORS["D"]
        ws6.cell(row=r, column=3, value=f"{f['weight_declared']:.2f}")
        ws6.cell(row=r, column=4, value=f"{f['weight_applied']:.3f}")
        ws6.cell(row=r, column=5, value=f["note"])
        for col in range(1, 6):
            ws6.cell(row=r, column=col).border = _THIN_BORDER
        r += 1

    r += 1
    ws6.cell(row=r, column=1,
             value="A factor with the same value in every ZIP cannot rank anything. It is "
                   "dropped and the remaining weights renormalized to 1.00, so scores span "
                   "the full 0-100. Left in, each one pulls every score toward 50 and "
                   "squeezes the reachable range — which is why no ZIP could earn an A.").font = _LABEL_FONT
    r += 2

    ws6.cell(row=r, column=1, value="Known limits of this run").font = _SUBTITLE_FONT
    r += 1
    for note in [
        "Distress density is a raw notice COUNT, not a density. It has no denominator "
        "(households or parcels per ZIP), so a large ZIP outranks a small one on size alone.",
        "Property values and equity come from whatever enrichment already sits in the CSVs. "
        "No Zillow call is made at analysis time.",
        "A valid in-state ZIP filed under the wrong county is not caught — only ZIPs outside "
        "the state's span are rejected.",
    ]:
        ws6.cell(row=r, column=1, value="\u2022  " + note).font = _LABEL_FONT
        r += 1

    if c.rejected_zips:
        r += 1
        ws6.cell(row=r, column=1, value="ZIPs rejected").font = _SUBTITLE_FONT
        r += 1
        for reason in c.rejected_zips.values():
            ws6.cell(row=r, column=1, value=reason).font = _LABEL_FONT
            r += 1

    r += 1
    ws6.cell(row=r, column=1, value="Files read").font = _SUBTITLE_FONT
    r += 1
    for name in c.files_read:
        ws6.cell(row=r, column=1, value=name).font = _LABEL_FONT
        r += 1

    ws6.column_dimensions["A"].width = 62
    ws6.column_dimensions["B"].width = 16
    ws6.column_dimensions["C"].width = 16
    ws6.column_dimensions["D"].width = 16
    ws6.column_dimensions["E"].width = 52

    # Save
    if not output_path:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = str(config.OUTPUT_DIR / f"market_analysis_{report.county}_{timestamp}.xlsx")

    wb.save(output_path)
    logger.info("Market report saved to %s", output_path)
    return output_path


# ── Main entry point ──────────────────────────────────────────────────

def run_market_analysis(counties: list[str] | None = None,
                        zip_codes: list[str] | None = None,
                        monthly_budget: float = 5000.0,
                        output_path: str = "") -> dict:
    """Run market analysis for specified counties.

    Returns dict with report data and output path.
    """
    counties = counties or ["Knox", "Blount"]
    county_str = ", ".join(counties)
    logger.info("Starting market analysis for: %s", county_str)

    # Step 1: Load and aggregate our own notice data
    corpus = CorpusStats()
    profiles = _load_notice_data(counties, stats=corpus)

    # If specific zips requested, filter
    if zip_codes:
        profiles = {k: v for k, v in profiles.items() if k in zip_codes}

    # Add known zips that may not have notices yet
    if not zip_codes:
        for county in counties:
            for z in COUNTY_ZIPS.get(county.lower(), []):
                if z not in profiles:
                    profiles[z] = ZipProfile(zip_code=z, county=county.title())

    if not profiles:
        logger.warning("No data found for analysis")
        return {"error": "No data found"}

    logger.info("Loaded notice data for %d zip codes", len(profiles))

    # Step 2: Score and rank
    factors: list = []
    scored = score_zip_codes(profiles, factor_log=factors)
    top = scored[:10]

    # Step 3: Budget allocation
    budget = _allocate_budget(scored, monthly_budget)

    # Step 4: Build report
    values = [z.median_value for z in scored if z.median_value > 0]
    equities = [z.avg_equity_pct for z in scored if z.avg_equity_pct > 0]

    report = MarketReport(
        county=county_str,
        analysis_date=datetime.now().strftime("%Y-%m-%d"),
        total_zips=len(scored),
        zip_profiles=scored,
        top_zips=top,
        total_notices=sum(z.total_notices for z in scored),
        avg_median_value=sum(values) / len(values) if values else 0,
        avg_equity=sum(equities) / len(equities) if equities else 0,
        total_distress=sum(z.total_notices for z in scored),
        corpus=corpus,
        factors=factors,
    )

    # Step 5: Generate Excel report
    report_path = generate_market_report(report, budget, output_path)

    logger.info("Market analysis complete: %d zips scored, top zip %s (score %.1f, grade %s)",
                len(scored), top[0].zip_code if top else "N/A",
                top[0].score if top else 0, top[0].grade if top else "N/A")

    return {
        "report": report,
        "budget": budget,
        "report_path": report_path,
    }
