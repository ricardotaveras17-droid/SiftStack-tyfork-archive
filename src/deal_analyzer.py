"""Deal analysis combining comping + rehab estimation into full deal packages.

Calculates MAO (Maximum Allowable Offer), ROI projections, holding costs,
selling costs, and financing scenarios for flip/wholesale/hold strategies.

Usage:
  python src/main.py analyze-deal --address "123 Main St, Knoxville, TN 37918"
  python src/main.py analyze-deal --address "123 Main St" --purchase-price 150000 --rehab-tier 2
"""

import logging
import math
from dataclasses import dataclass, field
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import config
from comp_analyzer import (
    ARVResult, SubjectProperty, fetch_comparable_sales, fetch_subject_property,
    calculate_arv, _fmt_money, DEFAULT_RADIUS_MILES, DEFAULT_MONTHS_BACK,
)
from rehab_estimator import (
    RehabEstimate, TIER_NAMES, estimate_rehab, estimate_wholetail,
    DEFAULT_REGION,
)

logger = logging.getLogger(__name__)

# ── Deal defaults ─────────────────────────────────────────────────────
DEFAULT_HARD_MONEY_RATE = 0.12     # 12% annual
DEFAULT_HARD_MONEY_POINTS = 2      # 2 origination points
DEFAULT_HARD_MONEY_TERM = 12       # months
DEFAULT_CONVENTIONAL_RATE = 0.07   # 7% annual (2026 rates)
DEFAULT_CONVENTIONAL_DOWN = 0.20   # 20% down
DEFAULT_INSURANCE_MONTHLY = 150.0
DEFAULT_UTILITIES_MONTHLY = 200.0
DEFAULT_AGENT_COMMISSION = 0.06    # 6% (3% buyer + 3% seller)
DEFAULT_CLOSING_COSTS_PCT = 0.025  # 2.5% of sale price
# Transfer tax is state-specific and, in New Jersey, steeply progressive. It
# lands on the RESALE side of a flip, so it is our cost. This module used to
# apply Tennessee's flat 0.37% to every deal in every state, which understated
# an NJ resale by roughly $1,700 on a $400K sale.
DEFAULT_TRANSFER_TAX_PCT = 0.0037  # Tennessee only: $0.37 per $100
DEFAULT_WHOLESALE_FEE = 10000.0
DEFAULT_FLIP_RULE = 0.75           # 75% Rule: MAO = ARV × 0.75 - rehab
DEFAULT_WHOLESALE_RULE = 0.70      # 70% Rule for wholesale
DEFAULT_CAP_RATE_TARGET = 0.08     # 8% target cap rate for hold
DEFAULT_CASH_ON_CASH_TARGET = 0.10 # 10% target CoC return
DEFAULT_VACANCY_RATE = 0.08        # 8% vacancy
DEFAULT_MAINTENANCE_PCT = 0.01     # 1% of value annually
DEFAULT_PROP_MGMT_PCT = 0.10       # 10% of rent

# ── Transfer tax ──────────────────────────────────────────────────────


class UnknownTransferTaxError(ValueError):
    """Raised when no transfer-tax schedule is calibrated for a state.

    Deliberately fatal, for the same reason the region lookup is: quietly
    applying one state's rate to another is a silent, systematic understatement
    of cost that flows straight into profit.
    """


# New Jersey Realty Transfer Fee. Seller-paid, graduated by portion.
# Source: NJ Division of Taxation — https://www.nj.gov/treasury/taxation/realty.shtml
# Verified 2026-08-20.
#
# NJ uses TWO bracket tables and switches wholly between them at $350,000; it is
# not one table with an extra bracket on top. A $350,001 sale is charged on the
# higher table from the first dollar, so crossing $350K costs about $630 more in
# fee than stopping just under it. That cliff is real and is worth knowing when
# pricing a resale near the line.
_NJ_RTF_TO_350K = (          # (bracket upper bound, dollars per $500)
    (150_000, 2.00),
    (200_000, 3.35),
    (350_000, 3.90),
)
_NJ_RTF_OVER_350K = (
    (150_000, 2.90),
    (200_000, 4.25),
    (550_000, 4.80),
    (850_000, 5.30),
    (1_000_000, 5.80),
    (float("inf"), 6.05),
)

# Graduated Percent Fee — replaced the buyer-paid 1% "mansion tax" for contracts
# fully executed on or after 2025-07-10, and is now SELLER-paid. Applies to
# Class 2 residential, Class 4A commercial, Class 4C co-ops and certain farms.
# These are NOT marginal rates: the tier rate applies to the ENTIRE
# consideration, so $2.2M is taxed 2% on all $2.2M ($44,000), not on the excess.
_NJ_GRADUATED_PCT = (        # (tier upper bound, rate on full consideration)
    (2_000_000, 0.010),
    (2_500_000, 0.020),
    (3_000_000, 0.025),
    (3_500_000, 0.030),
    (float("inf"), 0.035),
)

# Not modelled: NJ partial exemptions (senior citizen, blind, disabled,
# low/moderate-income housing) and the new-construction schedule. All of them
# REDUCE the fee, so ignoring them is conservative for underwriting.


def _nj_realty_transfer_fee(consideration: float) -> float:
    """NJ RTF plus, above $1M, the Graduated Percent Fee. Seller side."""
    if consideration <= 0:
        return 0.0

    # The statute charges each $500 "or fractional part thereof", so the
    # consideration rounds up to the next $500 before the brackets apply.
    taxable_total = math.ceil(consideration / 500.0) * 500.0

    table = _NJ_RTF_TO_350K if taxable_total <= 350_000 else _NJ_RTF_OVER_350K
    fee = 0.0
    lower = 0.0
    for upper, per_500 in table:
        if taxable_total <= lower:
            break
        in_bracket = min(taxable_total, upper) - lower
        fee += (in_bracket / 500.0) * per_500
        lower = upper

    if taxable_total > 1_000_000:
        for tier_upper, pct in _NJ_GRADUATED_PCT:
            if taxable_total <= tier_upper:
                fee += taxable_total * pct
                break

    return fee


# The property API returns the state inconsistently ("NJ" or "New Jersey"),
# so normalise before matching rather than raising on a spelling.
_STATE_ALIASES = {"NEW JERSEY": "NJ", "N.J.": "NJ", "TENNESSEE": "TN", "TENN.": "TN"}


def calculate_transfer_tax(sale_price: float, state: str = "") -> float:
    """Transfer tax on a sale, by state. Raises on an uncalibrated state."""
    st = (state or "").strip().upper()
    st = _STATE_ALIASES.get(st, st)
    if st == "NJ":
        return _nj_realty_transfer_fee(sale_price)
    if st == "TN":
        return sale_price * DEFAULT_TRANSFER_TAX_PCT
    raise UnknownTransferTaxError(
        f"No transfer-tax schedule is calibrated for state {state!r}. "
        "Calibrated states: NJ, TN. Add the state's schedule rather than "
        "borrowing another state's rate — transfer tax varies by an order of "
        "magnitude between states and lands directly in net profit."
    )


# ── Data structures ───────────────────────────────────────────────────


@dataclass
class MAOResult:
    """Maximum Allowable Offer calculations."""
    flip_mao: float = 0.0
    wholesale_mao: float = 0.0
    hold_mao: float = 0.0
    flip_rule_pct: float = DEFAULT_FLIP_RULE
    wholesale_rule_pct: float = DEFAULT_WHOLESALE_RULE


@dataclass
class HoldingCosts:
    """Monthly and total holding costs."""
    months: float = 0.0
    mortgage_monthly: float = 0.0
    taxes_monthly: float = 0.0
    insurance_monthly: float = DEFAULT_INSURANCE_MONTHLY
    utilities_monthly: float = DEFAULT_UTILITIES_MONTHLY
    total_monthly: float = 0.0
    total: float = 0.0


@dataclass
class SellingCosts:
    """Costs to sell the property."""
    agent_commission: float = 0.0
    closing_costs: float = 0.0
    transfer_tax: float = 0.0
    total: float = 0.0


@dataclass
class FlipProjection:
    """Profit projection for a flip."""
    arv: float = 0.0
    purchase_price: float = 0.0
    rehab_cost: float = 0.0
    holding_costs: float = 0.0
    selling_costs: float = 0.0
    total_investment: float = 0.0
    net_profit: float = 0.0
    roi_pct: float = 0.0
    months_to_complete: float = 0.0


@dataclass
class WholesaleProjection:
    """Profit projection for a wholesale assignment."""
    arv: float = 0.0
    contract_price: float = 0.0
    assignment_fee: float = 0.0
    buyer_rehab_cost: float = 0.0
    buyer_profit_estimate: float = 0.0


@dataclass
class HoldProjection:
    """Cash flow projection for buy-and-hold."""
    purchase_price: float = 0.0
    rehab_cost: float = 0.0
    total_investment: float = 0.0
    estimated_rent_monthly: float = 0.0
    gross_annual_income: float = 0.0
    vacancy_loss: float = 0.0
    effective_income: float = 0.0
    expenses_annual: float = 0.0
    noi: float = 0.0
    cap_rate: float = 0.0
    debt_service_annual: float = 0.0
    cash_flow_annual: float = 0.0
    cash_on_cash: float = 0.0
    equity_year5: float = 0.0


@dataclass
class FinancingScenario:
    """Financing comparison."""
    name: str = ""
    down_payment: float = 0.0
    loan_amount: float = 0.0
    rate: float = 0.0
    monthly_payment: float = 0.0
    total_cost_of_money: float = 0.0
    points_cost: float = 0.0


@dataclass
class DealPackage:
    """Complete deal analysis package."""
    subject: SubjectProperty = field(default_factory=SubjectProperty)
    arv: ARVResult = field(default_factory=ARVResult)
    rehab_full: RehabEstimate = field(default_factory=RehabEstimate)
    rehab_wholetail: RehabEstimate = field(default_factory=RehabEstimate)
    mao: MAOResult = field(default_factory=MAOResult)
    flip: FlipProjection = field(default_factory=FlipProjection)
    wholesale: WholesaleProjection = field(default_factory=WholesaleProjection)
    hold: HoldProjection = field(default_factory=HoldProjection)
    holding_costs: HoldingCosts = field(default_factory=HoldingCosts)
    selling_costs: SellingCosts = field(default_factory=SellingCosts)
    financing: list = field(default_factory=list)
    recommendation: str = ""
    risk_factors: list = field(default_factory=list)


# ── Calculation engine ────────────────────────────────────────────────

def _calc_monthly_payment(principal: float, annual_rate: float, months: int) -> float:
    """Standard amortization monthly payment."""
    if annual_rate == 0:
        return principal / months if months else 0
    r = annual_rate / 12
    return principal * (r * (1 + r) ** months) / ((1 + r) ** months - 1)


def _estimate_monthly_rent(arv: float, sqft: int, bedrooms: int) -> float:
    """Estimate monthly rent based on 1% rule and property characteristics."""
    # 1% rule as baseline (conservative for Knoxville)
    rent_pct_rule = arv * 0.008  # 0.8% for Knoxville (below 1% rule)
    # Sqft-based estimate
    rent_sqft = sqft * 0.75 if sqft else 0  # ~$0.75/sqft for Knoxville
    # Bedroom-based estimate
    rent_bed = 600 + (bedrooms * 200)  # base $600 + $200/bed
    # Average the estimates
    estimates = [e for e in [rent_pct_rule, rent_sqft, rent_bed] if e > 0]
    return round(sum(estimates) / len(estimates)) if estimates else 1200


def calculate_mao(arv_mid: float, rehab_cost: float,
                  wholesale_fee: float = DEFAULT_WHOLESALE_FEE) -> MAOResult:
    """Calculate Maximum Allowable Offer for each exit strategy."""
    flip_mao = (arv_mid * DEFAULT_FLIP_RULE) - rehab_cost
    wholesale_mao = (arv_mid * DEFAULT_WHOLESALE_RULE) - rehab_cost - wholesale_fee
    # Hold MAO: based on target cap rate and estimated rent
    # Simplified: 70% of ARV for hold (similar to wholesale)
    hold_mao = arv_mid * 0.70 - rehab_cost

    return MAOResult(
        flip_mao=max(0, round(flip_mao)),
        wholesale_mao=max(0, round(wholesale_mao)),
        hold_mao=max(0, round(hold_mao)),
    )


def calculate_holding_costs(purchase_price: float, rehab_months: float,
                            annual_tax: float = 0,
                            hard_money_rate: float = DEFAULT_HARD_MONEY_RATE) -> HoldingCosts:
    """Calculate holding costs during rehab period."""
    # Hard money interest-only payments
    mortgage_monthly = purchase_price * hard_money_rate / 12
    # Property taxes (use provided or estimate at 1% of purchase)
    taxes_monthly = (annual_tax / 12) if annual_tax else (purchase_price * 0.01 / 12)
    total_monthly = mortgage_monthly + taxes_monthly + DEFAULT_INSURANCE_MONTHLY + DEFAULT_UTILITIES_MONTHLY
    # Add 1 month for listing/closing after rehab
    total_months = rehab_months + 1

    return HoldingCosts(
        months=total_months,
        mortgage_monthly=round(mortgage_monthly),
        taxes_monthly=round(taxes_monthly),
        total_monthly=round(total_monthly),
        total=round(total_monthly * total_months),
    )


def calculate_selling_costs(sale_price: float, state: str = "") -> SellingCosts:
    """Calculate costs to sell the property.

    Raises:
        UnknownTransferTaxError: no transfer-tax schedule for ``state``.
    """
    commission = sale_price * DEFAULT_AGENT_COMMISSION
    closing = sale_price * DEFAULT_CLOSING_COSTS_PCT
    transfer = calculate_transfer_tax(sale_price, state)

    return SellingCosts(
        agent_commission=round(commission),
        closing_costs=round(closing),
        transfer_tax=round(transfer),
        total=round(commission + closing + transfer),
    )


def calculate_flip(arv_mid: float, purchase_price: float, rehab_cost: float,
                   holding: HoldingCosts, selling: SellingCosts,
                   rehab_months: float) -> FlipProjection:
    """Calculate flip profit projection."""
    total_investment = purchase_price + rehab_cost + holding.total + selling.total
    net_profit = arv_mid - total_investment
    roi = (net_profit / (purchase_price + rehab_cost) * 100) if (purchase_price + rehab_cost) else 0

    return FlipProjection(
        arv=round(arv_mid),
        purchase_price=round(purchase_price),
        rehab_cost=round(rehab_cost),
        holding_costs=round(holding.total),
        selling_costs=round(selling.total),
        total_investment=round(total_investment),
        net_profit=round(net_profit),
        roi_pct=round(roi, 1),
        months_to_complete=round(rehab_months + 2, 1),  # +2 for listing/closing
    )


def calculate_wholesale(arv_mid: float, contract_price: float,
                        rehab_cost: float,
                        assignment_fee: float = DEFAULT_WHOLESALE_FEE) -> WholesaleProjection:
    """Calculate wholesale assignment projection."""
    buyer_total = contract_price + assignment_fee + rehab_cost
    buyer_profit = arv_mid - buyer_total - (arv_mid * (DEFAULT_AGENT_COMMISSION + DEFAULT_CLOSING_COSTS_PCT))

    return WholesaleProjection(
        arv=round(arv_mid),
        contract_price=round(contract_price),
        assignment_fee=round(assignment_fee),
        buyer_rehab_cost=round(rehab_cost),
        buyer_profit_estimate=round(buyer_profit),
    )


def calculate_hold(purchase_price: float, rehab_cost: float, arv_mid: float,
                   sqft: int, bedrooms: int, annual_tax: float = 0) -> HoldProjection:
    """Calculate buy-and-hold cash flow projection."""
    total_investment = purchase_price + rehab_cost
    rent = _estimate_monthly_rent(arv_mid, sqft, bedrooms)
    gross_annual = rent * 12
    vacancy = gross_annual * DEFAULT_VACANCY_RATE
    effective = gross_annual - vacancy

    # Annual expenses
    taxes = annual_tax if annual_tax else purchase_price * 0.01
    insurance = DEFAULT_INSURANCE_MONTHLY * 12
    maintenance = arv_mid * DEFAULT_MAINTENANCE_PCT
    management = effective * DEFAULT_PROP_MGMT_PCT
    expenses = taxes + insurance + maintenance + management

    noi = effective - expenses
    cap_rate = (noi / total_investment * 100) if total_investment else 0

    # Debt service (conventional financing)
    down = total_investment * DEFAULT_CONVENTIONAL_DOWN
    loan = total_investment - down
    monthly_pmt = _calc_monthly_payment(loan, DEFAULT_CONVENTIONAL_RATE, 360)
    debt_service = monthly_pmt * 12

    cash_flow = noi - debt_service
    cash_on_cash = (cash_flow / down * 100) if down else 0

    # 5-year equity buildup (3% annual appreciation)
    equity_5yr = arv_mid * (1.03 ** 5) - arv_mid + down

    return HoldProjection(
        purchase_price=round(purchase_price),
        rehab_cost=round(rehab_cost),
        total_investment=round(total_investment),
        estimated_rent_monthly=round(rent),
        gross_annual_income=round(gross_annual),
        vacancy_loss=round(vacancy),
        effective_income=round(effective),
        expenses_annual=round(expenses),
        noi=round(noi),
        cap_rate=round(cap_rate, 1),
        debt_service_annual=round(debt_service),
        cash_flow_annual=round(cash_flow),
        cash_on_cash=round(cash_on_cash, 1),
        equity_year5=round(equity_5yr),
    )


def calculate_financing(purchase_price: float, rehab_cost: float) -> list[FinancingScenario]:
    """Calculate financing comparison scenarios."""
    total_needed = purchase_price + rehab_cost
    scenarios = []

    # Cash purchase
    scenarios.append(FinancingScenario(
        name="All Cash",
        down_payment=total_needed,
        loan_amount=0,
        rate=0,
        monthly_payment=0,
        total_cost_of_money=0,
        points_cost=0,
    ))

    # Hard money
    hm_points = total_needed * (DEFAULT_HARD_MONEY_POINTS / 100)
    hm_monthly = total_needed * DEFAULT_HARD_MONEY_RATE / 12
    scenarios.append(FinancingScenario(
        name="Hard Money (12%, 2pts, 12mo)",
        down_payment=hm_points,
        loan_amount=total_needed,
        rate=DEFAULT_HARD_MONEY_RATE,
        monthly_payment=round(hm_monthly),
        total_cost_of_money=round(hm_monthly * DEFAULT_HARD_MONEY_TERM + hm_points),
        points_cost=round(hm_points),
    ))

    # Conventional
    conv_down = total_needed * DEFAULT_CONVENTIONAL_DOWN
    conv_loan = total_needed - conv_down
    conv_monthly = _calc_monthly_payment(conv_loan, DEFAULT_CONVENTIONAL_RATE, 360)
    scenarios.append(FinancingScenario(
        name=f"Conventional ({DEFAULT_CONVENTIONAL_DOWN:.0%} down, {DEFAULT_CONVENTIONAL_RATE:.0%})",
        down_payment=round(conv_down),
        loan_amount=round(conv_loan),
        rate=DEFAULT_CONVENTIONAL_RATE,
        monthly_payment=round(conv_monthly),
        total_cost_of_money=round(conv_monthly * 360 - conv_loan),
        points_cost=0,
    ))

    # Seller financing (hypothetical)
    sf_down = total_needed * 0.10
    sf_loan = total_needed - sf_down
    sf_rate = 0.08
    sf_monthly = _calc_monthly_payment(sf_loan, sf_rate, 240)  # 20-year term
    scenarios.append(FinancingScenario(
        name="Seller Financing (10% down, 8%, 20yr)",
        down_payment=round(sf_down),
        loan_amount=round(sf_loan),
        rate=sf_rate,
        monthly_payment=round(sf_monthly),
        total_cost_of_money=round(sf_monthly * 240 - sf_loan),
        points_cost=0,
    ))

    return scenarios


def _assess_risk(arv: ARVResult, flip: FlipProjection, subject: SubjectProperty) -> list[str]:
    """Assess risk factors for the deal."""
    risks = []

    if arv.confidence == "low":
        risks.append(f"LOW CONFIDENCE ARV — {arv.confidence_reason}")
    if arv.spread_pct > 15:
        risks.append(f"Wide comp spread ({arv.spread_pct:.0f}%) — ARV uncertainty")
    if arv.comp_count < 3:
        risks.append(f"Only {arv.comp_count} comps — limited market data")
    if flip.roi_pct < 15:
        risks.append(f"Thin flip margin ({flip.roi_pct:.0f}% ROI) — little room for error")
    if flip.net_profit < 20000:
        risks.append(f"Low profit (${flip.net_profit:,.0f}) — may not justify effort/risk")
    if subject.year_built and subject.year_built < 1960:
        risks.append(f"Pre-1960 construction — potential lead paint, asbestos, foundation issues")
    if subject.mls_status and "active" in subject.mls_status.lower():
        risks.append("Property currently listed — competing with retail buyers")

    if not risks:
        risks.append("No major risk factors identified")

    return risks


def _make_recommendation(flip: FlipProjection, wholesale: WholesaleProjection,
                         hold: HoldProjection, arv: ARVResult) -> str:
    """Generate a go/no-go recommendation."""
    if arv.confidence == "none":
        return "INSUFFICIENT DATA — Cannot recommend. Need more comparable sales."

    strategies = []
    if flip.roi_pct >= 25 and flip.net_profit >= 25000:
        strategies.append(("FLIP", flip.roi_pct, f"${flip.net_profit:,.0f} profit, {flip.roi_pct:.0f}% ROI"))
    if wholesale.assignment_fee >= 5000 and wholesale.buyer_profit_estimate > 20000:
        strategies.append(("WHOLESALE", 100, f"${wholesale.assignment_fee:,.0f} assignment fee, buyer profits ${wholesale.buyer_profit_estimate:,.0f}"))
    if hold.cash_on_cash >= 8:
        strategies.append(("BUY & HOLD", hold.cash_on_cash, f"{hold.cash_on_cash:.1f}% CoC, ${hold.cash_flow_annual:,.0f}/yr cash flow"))

    if not strategies:
        if flip.roi_pct >= 15:
            return f"MARGINAL FLIP — {flip.roi_pct:.0f}% ROI. Proceed with caution; negotiate harder on price."
        return "NO-GO — Numbers don't work at this price point. Counter-offer or pass."

    strategies.sort(key=lambda x: x[1], reverse=True)
    best = strategies[0]
    return f"GO — Best strategy: {best[0]} ({best[2]})"


# ── Excel report generation ──────────────────────────────────────────

_HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
_HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
_TITLE_FONT = Font(name="Calibri", bold=True, size=16, color="2F5496")
_SUBTITLE_FONT = Font(name="Calibri", bold=True, size=12, color="333333")
_LABEL_FONT = Font(name="Calibri", size=11, color="555555")
_VALUE_FONT = Font(name="Calibri", bold=True, size=13, color="222222")
_GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
_GREEN_FONT = Font(name="Calibri", bold=True, color="006100")
_RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
_RED_FONT = Font(name="Calibri", bold=True, color="9C0006")
_YELLOW_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
_MONEY_FMT = '#,##0'
_THIN_BORDER = Border(bottom=Side(style="thin", color="D9D9D9"))


def _write_headers(ws, row, headers):
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN


def generate_deal_report(pkg: DealPackage, output_path: str = "") -> str:
    """Generate deal analysis Excel workbook."""
    wb = Workbook()

    # ── Tab 1: Deal Summary ───────────────────────────────────────
    ws = wb.active
    ws.title = "Deal Summary"
    ws.cell(row=1, column=1, value="Deal Analysis Report").font = _TITLE_FONT
    addr = f"{pkg.subject.address}, {pkg.subject.city}, {pkg.subject.state} {pkg.subject.zip_code}"
    ws.cell(row=2, column=1, value=addr).font = _SUBTITLE_FONT
    ws.cell(row=3, column=1, value=f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}").font = _LABEL_FONT

    # Recommendation
    rec_cell = ws.cell(row=5, column=1, value=pkg.recommendation)
    rec_cell.font = Font(name="Calibri", bold=True, size=14)
    if "GO —" in pkg.recommendation and "NO-GO" not in pkg.recommendation:
        rec_cell.fill = _GREEN_FILL
        rec_cell.font = Font(name="Calibri", bold=True, size=14, color="006100")
    elif "NO-GO" in pkg.recommendation:
        rec_cell.fill = _RED_FILL
        rec_cell.font = Font(name="Calibri", bold=True, size=14, color="9C0006")
    else:
        rec_cell.fill = _YELLOW_FILL

    row = 7
    data = [
        ("ARV (Recommended)", _fmt_money(pkg.arv.arv_mid)),
        ("ARV Confidence", pkg.arv.confidence.upper()),
        ("", ""),
        ("MAO — Flip (75% Rule)", _fmt_money(pkg.mao.flip_mao)),
        ("MAO — Wholesale (70% Rule)", _fmt_money(pkg.mao.wholesale_mao)),
        ("MAO — Hold", _fmt_money(pkg.mao.hold_mao)),
        ("", ""),
        ("Full Rehab Cost", _fmt_money(pkg.rehab_full.grand_total)),
        ("Wholetail Cost", _fmt_money(pkg.rehab_wholetail.grand_total)),
        ("", ""),
        ("Flip Profit", _fmt_money(pkg.flip.net_profit)),
        ("Flip ROI", f"{pkg.flip.roi_pct:.1f}%"),
        ("Flip Timeline", f"{pkg.flip.months_to_complete:.0f} months"),
        ("", ""),
        ("Wholesale Fee", _fmt_money(pkg.wholesale.assignment_fee)),
        ("", ""),
        ("Hold Cash Flow (Annual)", _fmt_money(pkg.hold.cash_flow_annual)),
        ("Hold Cash-on-Cash", f"{pkg.hold.cash_on_cash:.1f}%"),
        ("Hold Cap Rate", f"{pkg.hold.cap_rate:.1f}%"),
    ]
    for label, value in data:
        ws.cell(row=row, column=1, value=label).font = _LABEL_FONT
        cell = ws.cell(row=row, column=2, value=value)
        cell.font = _VALUE_FONT
        if "Profit" in label and pkg.flip.net_profit > 0:
            cell.fill = _GREEN_FILL
            cell.font = _GREEN_FONT
        elif "Profit" in label and pkg.flip.net_profit <= 0:
            cell.fill = _RED_FILL
            cell.font = _RED_FONT
        row += 1

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 25

    # ── Tab 2: ARV Analysis ───────────────────────────────────────
    ws2 = wb.create_sheet("ARV Analysis")
    ws2.cell(row=1, column=1, value="ARV Analysis (from Comp Report)").font = _TITLE_FONT
    arv_data = [
        ("ARV Low", _fmt_money(pkg.arv.arv_low)),
        ("ARV Mid (Recommended)", _fmt_money(pkg.arv.arv_mid)),
        ("ARV High", _fmt_money(pkg.arv.arv_high)),
        ("Confidence", pkg.arv.confidence.upper()),
        ("Reason", pkg.arv.confidence_reason),
        ("Comps Used", str(pkg.arv.comp_count)),
        ("Avg PPSF", f"${pkg.arv.ppsf_avg:,.2f}"),
        ("PPSF Range", f"${pkg.arv.ppsf_range[0]:,.2f} — ${pkg.arv.ppsf_range[1]:,.2f}"),
        ("Avg Adjustment", _fmt_money(pkg.arv.avg_adjustment)),
        ("Spread", f"{pkg.arv.spread_pct:.1f}%"),
    ]
    for i, (label, value) in enumerate(arv_data, 3):
        ws2.cell(row=i, column=1, value=label).font = _LABEL_FONT
        ws2.cell(row=i, column=2, value=value).font = _VALUE_FONT
    ws2.column_dimensions["A"].width = 25
    ws2.column_dimensions["B"].width = 30

    # ── Tab 3: Rehab Budget ───────────────────────────────────────
    ws3 = wb.create_sheet("Rehab Budget")
    ws3.cell(row=1, column=1, value="Rehab Budget Summary").font = _TITLE_FONT
    _write_headers(ws3, 3, ["Category", "Full Rehab", "Wholetail"])
    full_dict = {r.category: r.total for r in pkg.rehab_full.rooms}
    wt_dict = {r.category: r.total for r in pkg.rehab_wholetail.rooms}
    all_cats = list(dict.fromkeys([r.category for r in pkg.rehab_full.rooms] +
                                   [r.category for r in pkg.rehab_wholetail.rooms]))
    for i, cat in enumerate(all_cats, 4):
        ws3.cell(row=i, column=1, value=cat)
        ws3.cell(row=i, column=2, value=full_dict.get(cat, 0)).number_format = _MONEY_FMT
        ws3.cell(row=i, column=3, value=wt_dict.get(cat, 0)).number_format = _MONEY_FMT
        for c in range(1, 4):
            ws3.cell(row=i, column=c).border = _THIN_BORDER
    tot_row = 4 + len(all_cats)
    ws3.cell(row=tot_row, column=1, value="GRAND TOTAL").font = _VALUE_FONT
    ws3.cell(row=tot_row, column=2, value=pkg.rehab_full.grand_total).number_format = _MONEY_FMT
    ws3.cell(row=tot_row, column=2).font = _VALUE_FONT
    ws3.cell(row=tot_row, column=3, value=pkg.rehab_wholetail.grand_total).number_format = _MONEY_FMT
    ws3.column_dimensions["A"].width = 25
    ws3.column_dimensions["B"].width = 15
    ws3.column_dimensions["C"].width = 15

    # ── Tab 4: MAO Calculation ────────────────────────────────────
    ws4 = wb.create_sheet("MAO Calculation")
    ws4.cell(row=1, column=1, value="Maximum Allowable Offer").font = _TITLE_FONT
    mao_data = [
        ("FLIP (75% Rule)", "", ""),
        ("Formula", "ARV × 75% - Rehab", ""),
        ("ARV", _fmt_money(pkg.arv.arv_mid), ""),
        ("ARV × 75%", _fmt_money(pkg.arv.arv_mid * 0.75), ""),
        ("Less Rehab", f"- {_fmt_money(pkg.rehab_full.grand_total)}", ""),
        ("MAO", _fmt_money(pkg.mao.flip_mao), ""),
        ("", "", ""),
        ("WHOLESALE (70% Rule)", "", ""),
        ("Formula", "ARV × 70% - Rehab - Assignment Fee", ""),
        ("ARV × 70%", _fmt_money(pkg.arv.arv_mid * 0.70), ""),
        ("Less Rehab", f"- {_fmt_money(pkg.rehab_full.grand_total)}", ""),
        ("Less Assignment", f"- {_fmt_money(DEFAULT_WHOLESALE_FEE)}", ""),
        ("MAO", _fmt_money(pkg.mao.wholesale_mao), ""),
    ]
    for i, (label, value, _) in enumerate(mao_data, 3):
        ws4.cell(row=i, column=1, value=label).font = _LABEL_FONT if label else _SUBTITLE_FONT
        ws4.cell(row=i, column=2, value=value).font = _VALUE_FONT
        if label == "MAO":
            ws4.cell(row=i, column=2).fill = _GREEN_FILL
            ws4.cell(row=i, column=2).font = _GREEN_FONT
    ws4.column_dimensions["A"].width = 25
    ws4.column_dimensions["B"].width = 30

    # ── Tab 5: Flip Projection ────────────────────────────────────
    ws5 = wb.create_sheet("Flip Projection")
    ws5.cell(row=1, column=1, value="Flip Profit Projection").font = _TITLE_FONT
    flip_data = [
        ("Sale Price (ARV)", pkg.flip.arv),
        ("", 0),
        ("Less: Purchase Price", -pkg.flip.purchase_price),
        ("Less: Rehab Cost", -pkg.flip.rehab_cost),
        ("Less: Holding Costs", -pkg.flip.holding_costs),
        ("Less: Selling Costs", -pkg.flip.selling_costs),
        ("", 0),
        ("NET PROFIT", pkg.flip.net_profit),
        ("ROI", None),
        ("Timeline", None),
    ]
    for i, (label, value) in enumerate(flip_data, 3):
        ws5.cell(row=i, column=1, value=label).font = _LABEL_FONT
        if value is not None:
            ws5.cell(row=i, column=2, value=value).number_format = _MONEY_FMT
        if label == "NET PROFIT":
            ws5.cell(row=i, column=2).font = Font(name="Calibri", bold=True, size=14, color="006100" if value > 0 else "9C0006")
        elif label == "ROI":
            ws5.cell(row=i, column=2, value=f"{pkg.flip.roi_pct:.1f}%").font = _VALUE_FONT
        elif label == "Timeline":
            ws5.cell(row=i, column=2, value=f"{pkg.flip.months_to_complete:.0f} months").font = _VALUE_FONT
    ws5.column_dimensions["A"].width = 25
    ws5.column_dimensions["B"].width = 20

    # ── Tab 6: Hold Projection ────────────────────────────────────
    ws6 = wb.create_sheet("Hold Projection")
    ws6.cell(row=1, column=1, value="Buy & Hold Cash Flow Analysis").font = _TITLE_FONT
    hold_data = [
        ("Total Investment", _fmt_money(pkg.hold.total_investment)),
        ("Est. Monthly Rent", _fmt_money(pkg.hold.estimated_rent_monthly)),
        ("", ""),
        ("Gross Annual Income", _fmt_money(pkg.hold.gross_annual_income)),
        ("Less: Vacancy (8%)", f"- {_fmt_money(pkg.hold.vacancy_loss)}"),
        ("Effective Income", _fmt_money(pkg.hold.effective_income)),
        ("Less: Expenses", f"- {_fmt_money(pkg.hold.expenses_annual)}"),
        ("NOI", _fmt_money(pkg.hold.noi)),
        ("Less: Debt Service", f"- {_fmt_money(pkg.hold.debt_service_annual)}"),
        ("", ""),
        ("ANNUAL CASH FLOW", _fmt_money(pkg.hold.cash_flow_annual)),
        ("Cap Rate", f"{pkg.hold.cap_rate:.1f}%"),
        ("Cash-on-Cash Return", f"{pkg.hold.cash_on_cash:.1f}%"),
        ("5-Year Equity Buildup", _fmt_money(pkg.hold.equity_year5)),
    ]
    for i, (label, value) in enumerate(hold_data, 3):
        ws6.cell(row=i, column=1, value=label).font = _LABEL_FONT
        ws6.cell(row=i, column=2, value=value).font = _VALUE_FONT
    ws6.column_dimensions["A"].width = 25
    ws6.column_dimensions["B"].width = 20

    # ── Tab 7: Financing ──────────────────────────────────────────
    ws7 = wb.create_sheet("Financing")
    ws7.cell(row=1, column=1, value="Financing Comparison").font = _TITLE_FONT
    _write_headers(ws7, 3, ["Scenario", "Down Payment", "Loan Amount", "Rate",
                             "Monthly Payment", "Total Cost of Money"])
    for i, fin in enumerate(pkg.financing, 4):
        ws7.cell(row=i, column=1, value=fin.name)
        ws7.cell(row=i, column=2, value=fin.down_payment).number_format = _MONEY_FMT
        ws7.cell(row=i, column=3, value=fin.loan_amount).number_format = _MONEY_FMT
        ws7.cell(row=i, column=4, value=f"{fin.rate:.1%}" if fin.rate else "N/A")
        ws7.cell(row=i, column=5, value=fin.monthly_payment).number_format = _MONEY_FMT
        ws7.cell(row=i, column=6, value=fin.total_cost_of_money).number_format = _MONEY_FMT
        for c in range(1, 7):
            ws7.cell(row=i, column=c).border = _THIN_BORDER
    for col, w in enumerate([35, 15, 15, 10, 18, 22], 1):
        ws7.column_dimensions[chr(64 + col)].width = w

    # ── Tab 8: Risk Factors ───────────────────────────────────────
    ws8 = wb.create_sheet("Risk Factors")
    ws8.cell(row=1, column=1, value="Risk Assessment").font = _TITLE_FONT
    for i, risk in enumerate(pkg.risk_factors, 3):
        cell = ws8.cell(row=i, column=1, value=f"• {risk}")
        cell.font = _LABEL_FONT
        if "LOW" in risk or "Wide" in risk or "Only" in risk or "Thin" in risk:
            cell.font = Font(name="Calibri", color="9C0006")
    ws8.column_dimensions["A"].width = 70

    # Save
    if not output_path:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_addr = "".join(c if c.isalnum() or c in " -" else "_" for c in pkg.subject.address)[:40]
        output_path = str(config.OUTPUT_DIR / f"deal_analysis_{safe_addr}_{timestamp}.xlsx")

    wb.save(output_path)
    logger.info("Deal analysis saved to %s", output_path)
    return output_path


# ── Main entry point ──────────────────────────────────────────────────

def run_deal_analysis(address: str, city: str = "", state: str = "TN",
                      zip_code: str = "", purchase_price: float = 0,
                      rehab_tier: int = 2, exit_strategy: str = "flip",
                      region: str = DEFAULT_REGION, county: str = "",
                      walkthrough_verified: bool = False,
                      radius: float = DEFAULT_RADIUS_MILES,
                      months: int = DEFAULT_MONTHS_BACK,
                      output_path: str = "") -> dict:
    """Run complete deal analysis: comp → rehab → MAO → projections → report.

    Returns dict with the DealPackage and report path.
    """
    logger.info("Starting deal analysis for: %s", address)

    # Step 1: Fetch subject property
    subject = fetch_subject_property(address, city, state, zip_code)
    if not subject:
        return {"error": "Could not fetch property data"}

    # Step 2: Fetch comps and calculate ARV
    comps = fetch_comparable_sales(subject, radius, months)
    # walkthrough_verified gates the reconfigure-to-more-bedrooms upside. Left
    # False the upside is computed and shown but never priced into MAO.
    arv = calculate_arv(subject, comps, walkthrough_verified=walkthrough_verified)

    if arv.confidence == "none":
        # Fail loud. Continuing here produced a fully-rendered, branded deal report
        # in which ARV, MAO, and every projection were $0 — output that looks
        # deliberate rather than broken. A missing ARV is not a degraded result,
        # it is the absence of one.
        logger.error(
            "No ARV could be calculated from %d comp(s) — refusing to emit a deal "
            "package. Reason: %s", len(comps),
            arv.confidence_reason or "insufficient comp data")
        return {
            "error": "No ARV could be calculated — "
                     f"{arv.confidence_reason or 'insufficient comp data'}",
            "subject": subject,
            "comps": comps,
        }

    # Step 3: Rehab estimates
    rehab_full = estimate_rehab(
        address=subject.address, sqft=subject.sqft, bedrooms=subject.bedrooms,
        bathrooms=subject.bathrooms, year_built=subject.year_built,
        tier=rehab_tier, scope="full", region=region,
        state=subject.state, county=county, city=subject.city,
    )
    rehab_wholetail = estimate_wholetail(
        address=subject.address, sqft=subject.sqft, bedrooms=subject.bedrooms,
        bathrooms=subject.bathrooms, year_built=subject.year_built, region=region,
        state=subject.state, county=county, city=subject.city,
    )

    # Use purchase price or default to flip MAO
    if not purchase_price:
        purchase_price = (arv.arv_mid * DEFAULT_FLIP_RULE) - rehab_full.grand_total
        purchase_price = max(0, purchase_price)
        logger.info("No purchase price given — using flip MAO: %s", _fmt_money(purchase_price))

    # Step 4: MAO
    mao = calculate_mao(arv.arv_mid, rehab_full.grand_total)

    # Step 5: Projections
    rehab_months = rehab_full.total_weeks / 4.0  # weeks to months
    holding = calculate_holding_costs(purchase_price, rehab_months)
    selling = calculate_selling_costs(arv.arv_mid, state=subject.state)

    flip = calculate_flip(arv.arv_mid, purchase_price, rehab_full.grand_total,
                          holding, selling, rehab_months)
    wholesale = calculate_wholesale(arv.arv_mid, purchase_price, rehab_full.grand_total)
    hold = calculate_hold(purchase_price, rehab_full.grand_total, arv.arv_mid,
                          subject.sqft, subject.bedrooms)
    financing = calculate_financing(purchase_price, rehab_full.grand_total)

    # Step 6: Risk assessment and recommendation
    risk_factors = _assess_risk(arv, flip, subject)
    recommendation = _make_recommendation(flip, wholesale, hold, arv)

    # Assemble package
    pkg = DealPackage(
        subject=subject,
        arv=arv,
        rehab_full=rehab_full,
        rehab_wholetail=rehab_wholetail,
        mao=mao,
        flip=flip,
        wholesale=wholesale,
        hold=hold,
        holding_costs=holding,
        selling_costs=selling,
        financing=financing,
        recommendation=recommendation,
        risk_factors=risk_factors,
    )

    # Step 7: Generate report
    report_path = generate_deal_report(pkg, output_path)

    logger.info("Deal analysis complete: %s | ARV %s | Flip profit %s (%s%% ROI) | %s",
                subject.address, _fmt_money(arv.arv_mid),
                _fmt_money(flip.net_profit), f"{flip.roi_pct:.0f}",
                recommendation.split("—")[0].strip())

    return {
        "package": pkg,
        "report_path": report_path,
    }
