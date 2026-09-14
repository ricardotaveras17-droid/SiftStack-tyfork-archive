"""Lender package: the two safety properties, and formula/Python agreement.

    .venv/bin/python tests/test_lender_package.py
"""
import json, math, re, sys, tempfile
from pathlib import Path
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook                                    # noqa: E402
import lender_package_nj as lp                                        # noqa: E402
from deal_analyzer import calculate_transfer_tax                      # noqa: E402

SPEC = json.loads((ROOT / "deals" / "EXAMPLE_nj_lender_spec.json").read_text())
fails = 0


def check(label, ok, detail=""):
    global fails
    print(f"  {'OK  ' if ok else 'FAIL'} — {label}{(': ' + detail) if detail else ''}")
    if not ok:
        fails += 1


# ── 1. the live NJ transfer-fee formula must equal the Python ─────────
def eval_excel(expr: str, value: float) -> float:
    """Tiny evaluator for the subset of Excel used by the RTF formula."""
    py = expr.replace("ARVLow", "V").replace("ARV", "V")
    py = re.sub(r"\bMIN\(", "min(", py)
    py = re.sub(r"\bMAX\(", "max(", py)
    py = re.sub(r"\bCEILING\(([^,]+),\s*([0-9]+)\)", r"(math.ceil((\1)/\2)*\2)", py)
    # IF(a,b,c) -> (b if a else c), innermost first
    while "IF(" in py:
        i = py.rindex("IF(")
        depth, j = 0, i + 3
        while j < len(py):
            if py[j] == "(":
                depth += 1
            elif py[j] == ")":
                if depth == 0:
                    break
                depth -= 1
            j += 1
        args, depth, cur = [], 0, ""
        for ch in py[i + 3:j]:
            if ch == "," and depth == 0:
                args.append(cur); cur = ""
            else:
                if ch == "(": depth += 1
                elif ch == ")": depth -= 1
                cur += ch
        args.append(cur)
        py = py[:i] + f"(({args[1]}) if ({args[0]}) else ({args[2]}))" + py[j + 1:]
    return eval(py, {"math": math, "min": min, "max": max}, {"V": value})


print("== live Excel RTF formula vs calculate_transfer_tax ==")
formula = lp.nj_transfer_fee_formula("ARV")
for price in (150_000, 250_000, 350_000, 350_500, 400_000, 700_000, 1_000_000, 2_200_000):
    xl = eval_excel(formula, price)
    py = calculate_transfer_tax(price, "NJ")
    check(f"${price:,}", abs(xl - py) < 0.01, f"excel {xl:,.2f} vs python {py:,.2f}")

# ── 2. DayOne must be a FORMULA, never a typed literal ────────────────
print("\n== day-one advance is derived, not typed ==")
out = Path(tempfile.gettempdir()) / "lp_test.xlsx"
lp.build_workbook(SPEC, str(out), force=True)
wb = load_workbook(out, data_only=False)
ts = wb["Term Sheet"]
cells = {ts.cell(row=r, column=1).value: ts.cell(row=r, column=2).value
         for r in range(1, ts.max_row + 1)}
day_one = cells.get("Day one advance at closing")
check("cell holds a formula", isinstance(day_one, str) and day_one.startswith("="),
      repr(day_one))
check("derives from Loan and RehabTotal",
      isinstance(day_one, str) and "Loan" in day_one and "RehabTotal" in day_one)
hold = cells.get("Renovation holdback")
check("holdback also derived", isinstance(hold, str) and hold.startswith("="), repr(hold))

# ── 3. defined names exist so the inputs actually drive the sheet ─────
print("\n== defined names wired ==")
names = set(wb.defined_names)
missing = [n for n, *_ in lp.INPUTS if n not in names]
check(f"all {len(lp.INPUTS)} inputs have defined names", not missing, str(missing))

# ── 4. read-back protection ──────────────────────────────────────────
print("\n== read-back protection ==")
wb2 = load_workbook(out)
wb2["Term Sheet"]["B7"] = 999_000        # a reviewer edits the loan amount
wb2.save(out)
diffs = lp.diff_against_spec(lp.read_back_inputs(str(out)), SPEC)
check("edit is detected", any("Loan" in d for d in diffs), "; ".join(diffs))
try:
    lp.build_workbook(SPEC, str(out))
    check("rebuild refuses without --force", False)
except RuntimeError as exc:
    check("rebuild refuses without --force", True, str(exc).splitlines()[0][:60])
try:
    lp.build_workbook(SPEC, str(out), force=True)
    check("--force overwrites", True)
except RuntimeError:
    check("--force overwrites", False)

# ── 5. compute() sanity ──────────────────────────────────────────────
print("\n== python model ==")
m = lp.compute(SPEC)
check("day_one = loan - rehab", m["day_one"] == m["loan"] - m["rehab"])
check("holdback ties to rehab budget", m["rehab_tranche"] == m["rehab"],
      f"{m['rehab_tranche']:,} vs {m['rehab']:,}")
check("per draw x draws = holdback",
      abs(m["per_draw"] * m["draws"] - m["rehab_tranche"]) <= m["draws"])
out.unlink(missing_ok=True)

print(f"\n{'FAILURES: ' + str(fails) if fails else 'all checks passed'}")
sys.exit(1 if fails else 0)
